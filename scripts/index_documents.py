"""
Document indexing script for TEIA RAG system.

This script:
1. Loads documents from the configured data path (PDFs, Excel, Markdown, Text)
2. Chunks them into smaller pieces
3. Generates embeddings
4. Stores in ChromaDB for retrieval

Usage:
    python scripts/index_documents.py

Run this script whenever you add new course materials.
"""
import sys
from pathlib import Path

# Add src to path for imports
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

import json
from typing import List, Dict, Any

import chromadb
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from openpyxl import load_workbook

from config import settings
from services.embedding_service import get_embedding_service
from utils.logger import get_logger

logger = get_logger("indexer")

# Chunking configuration
CHUNK_SIZE = 500  # characters
CHUNK_OVERLAP = 50  # characters

# Module mapping based on filename patterns
MODULE_PATTERNS = {
    "caracterizacion": ["caracterizacion", "caracterización", "asignatura"],
    "factores_situacionales": ["factores", "situacional", "contexto"],
    "resultados_aprendizaje": ["resultado", "aprendizaje", "learning", "outcome"],
    "actividades_aprendizaje": ["actividad", "activity", "tarea"],
    "evaluacion": ["evaluacion", "evaluación", "assessment", "rubrica", "rúbrica"],
    "secuencia": ["secuencia", "sequence", "cronograma", "schedule"],
}


def detect_module(filename: str, content: str = "") -> str:
    """Detect module based on filename or content patterns."""
    filename_lower = filename.lower()
    content_lower = content.lower()[:500]  # Check first 500 chars

    for module, patterns in MODULE_PATTERNS.items():
        for pattern in patterns:
            if pattern in filename_lower or pattern in content_lower:
                return module

    return "general"


def load_pdf(file_path: Path) -> str:
    """Extract text from PDF file."""
    reader = PdfReader(file_path)
    text_parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            text_parts.append(text)
    return "\n\n".join(text_parts)


def load_excel(file_path: Path) -> str:
    """Extract text from Excel file."""
    workbook = load_workbook(file_path, data_only=True)
    text_parts = []

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        text_parts.append(f"=== Hoja: {sheet_name} ===\n")

        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
            if row_text.strip():
                text_parts.append(row_text)

    return "\n".join(text_parts)


def load_markdown(file_path: Path) -> str:
    """Load markdown file."""
    return file_path.read_text(encoding="utf-8")


def load_text(file_path: Path) -> str:
    """Load plain text file."""
    return file_path.read_text(encoding="utf-8")


def load_document(file_path: Path) -> Dict[str, Any]:
    """
    Load a document and return its content with metadata.

    Returns:
        Dict with 'content', 'source', 'file_type', 'module'
    """
    suffix = file_path.suffix.lower()

    loaders = {
        ".pdf": load_pdf,
        ".xlsx": load_excel,
        ".xls": load_excel,
        ".md": load_markdown,
        ".txt": load_text,
    }

    loader = loaders.get(suffix)
    if not loader:
        logger.debug(f"Skipping unsupported file type: {file_path.name}")
        return None

    try:
        content = loader(file_path)
        module = detect_module(file_path.name, content)

        return {
            "content": content,
            "source": file_path.name,
            "file_type": suffix[1:],  # Remove the dot
            "module": module,
            "path": str(file_path.relative_to(settings.DATA_RAW_PATH)),
        }
    except Exception as e:
        logger.error(f"Error loading {file_path.name}: {e}")
        return None


def chunk_document(doc: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Split document into chunks with metadata."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    chunks = splitter.split_text(doc["content"])

    return [
        {
            "content": chunk,
            "source": doc["source"],
            "file_type": doc["file_type"],
            "module": doc["module"],
            "chunk_index": i,
            "total_chunks": len(chunks),
        }
        for i, chunk in enumerate(chunks)
    ]


def get_all_documents() -> List[Path]:
    """
    Get all document files from the configured raw data path.
    Scans recursively and filters by supported extensions.
    """
    documents = []
    raw_path = settings.DATA_RAW_PATH
    supported_extensions = set(settings.SUPPORTED_EXTENSIONS)

    if not raw_path.exists():
        logger.warning(f"Raw data path does not exist: {raw_path}")
        return documents

    # Scan recursively for all supported file types
    for file_path in raw_path.rglob("*"):
        if file_path.is_file() and not file_path.name.startswith("."):
            if file_path.suffix.lower() in supported_extensions:
                documents.append(file_path)

    return documents


def save_processed_chunks(chunks: List[Dict[str, Any]]):
    """Save processed chunks to JSON for debugging/review."""
    processed_path = settings.DATA_PROCESSED_PATH
    processed_path.mkdir(parents=True, exist_ok=True)
    output_file = processed_path / "chunks.json"
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(chunks, f, ensure_ascii=False, indent=2)
    logger.info(f"Saved {len(chunks)} chunks to {output_file}")


def index_documents():
    """Main function to index all documents."""
    logger.info("=" * 60)
    logger.info("TEIA Document Indexer")
    logger.info("=" * 60)

    # Get embedding service
    logger.info("Loading embedding model...")
    embedding_service = get_embedding_service()
    logger.info(f"Embedding dimension: {embedding_service.dimension}")

    # Initialize ChromaDB
    logger.info("Initializing ChromaDB...")
    chroma_path = settings.CHROMA_DB_PATH
    chroma_path.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(chroma_path))

    # Delete existing collection if it exists (fresh index)
    try:
        client.delete_collection("teia_course_content")
        logger.info("Deleted existing collection")
    except Exception as e:
        logger.debug(f"No existing collection to delete: {e}")

    collection = client.create_collection(
        name="teia_course_content",
        metadata={"description": "TEIA course content for RAG retrieval"}
    )
    logger.info("Created new collection: teia_course_content")

    # Load and process documents
    logger.info("Scanning for documents...")
    document_paths = get_all_documents()
    logger.info(f"Found {len(document_paths)} documents")

    if not document_paths:
        logger.warning("No documents found! Please add documents to:")
        logger.warning(f"  {settings.DATA_RAW_PATH}")
        logger.warning(f"Supported extensions: {', '.join(settings.SUPPORTED_EXTENSIONS)}")
        return

    all_chunks = []

    for doc_path in document_paths:
        logger.info(f"Processing: {doc_path.name}")

        doc = load_document(doc_path)
        if doc is None:
            continue

        chunks = chunk_document(doc)
        logger.debug(f"  Module: {doc['module']}, Chunks: {len(chunks)}")

        all_chunks.extend(chunks)

    logger.info(f"Total chunks to index: {len(all_chunks)}")

    if not all_chunks:
        logger.warning("No chunks to index!")
        return

    # Save processed chunks for review
    save_processed_chunks(all_chunks)

    # Generate embeddings
    logger.info("Generating embeddings...")
    texts = [chunk["content"] for chunk in all_chunks]
    embeddings = embedding_service.embed_texts(texts)
    logger.info(f"Generated {len(embeddings)} embeddings")

    # Add to ChromaDB
    logger.info("Adding to ChromaDB...")
    collection.add(
        ids=[f"chunk_{i}" for i in range(len(all_chunks))],
        embeddings=embeddings,
        documents=texts,
        metadatas=[
            {
                "source": chunk["source"],
                "module": chunk["module"],
                "file_type": chunk["file_type"],
                "chunk_index": chunk["chunk_index"],
                "total_chunks": chunk["total_chunks"],
            }
            for chunk in all_chunks
        ],
    )

    logger.info(f"Successfully indexed {len(all_chunks)} chunks!")
    logger.info(f"ChromaDB path: {chroma_path}")

    # Summary by module
    logger.info("Chunks by module:")
    module_counts = {}
    for chunk in all_chunks:
        module = chunk["module"]
        module_counts[module] = module_counts.get(module, 0) + 1

    for module, count in sorted(module_counts.items()):
        logger.info(f"  {module}: {count} chunks")

    logger.info("=" * 60)
    logger.info("Indexing complete!")
    logger.info("=" * 60)


if __name__ == "__main__":
    index_documents()
